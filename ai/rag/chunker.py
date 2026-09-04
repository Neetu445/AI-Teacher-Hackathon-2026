import os
import re
from pypdf import PdfReader

def extract_text_from_pdf(file_path: str) -> str:
    """Backward compat: returns concatenated text."""
    pages = extract_pdf_pages(file_path)
    return _clean_text("\n".join([p["text"] for p in pages]).strip())

# Track whether last PDF extraction used OCR (for adaptive threshold)
_last_ocr_used: bool = False

def extract_pdf_pages(file_path: str) -> list[dict]:
    """Extract per-page text with page numbers. Returns [{'page':1,'text':...},...]"""
    global _last_ocr_used
    _last_ocr_used = False
    pages = []
    try:
        reader = PdfReader(file_path)
        for idx, page in enumerate(reader.pages):
            try:
                t = _clean_text((page.extract_text() or "").strip())
            except Exception:
                t = ""
            if t:
                pages.append({"page": idx + 1, "text": t})
    except Exception as e:
        print(f"pypdf failed, will try OCR: {e}")
        pages = []
    total_chars = sum(len(p["text"]) for p in pages)
    # Fallback to OCR if pypdf extracted almost nothing (scanned PDF)
    if total_chars < 100 or (total_chars < 300 and all(p["text"].replace(" ", "").replace("\n","").strip().isdigit() for p in pages if p["text"])):
        try:
            print("Scanned PDF detected - running OCR (may take 20-30s first time)...")
            ocr_pages = _ocr_pdf_pages(file_path)
            if ocr_pages and sum(len(p["text"]) for p in ocr_pages) > total_chars:
                print(f"OCR success: {sum(len(p['text']) for p in ocr_pages)} chars across {len(ocr_pages)} pages")
                _last_ocr_used = True
                return ocr_pages
            else:
                print("OCR returned less text than pypdf, using pypdf")
        except Exception as e:
            print(f"OCR fallback failed: {e}")
    return pages

def was_last_ocr_used() -> bool:
    """Return whether last extract_pdf_pages() used OCR fallback."""
    return _last_ocr_used

_ocr_reader = None
def _ocr_pdf(file_path: str) -> str:
    global _ocr_reader
    try:
        import fitz  # PyMuPDF
    except ImportError:
        print("OCR skipped: PyMuPDF not installed. Run: pip install PyMuPDF")
        return ""
    try:
        import easyocr
        import numpy as np
        import cv2
    except ImportError as e:
        print(f"OCR skipped: missing dep {e}. Run: pip install easyocr opencv-python-headless")
        return ""
    # init easyocr once
    if _ocr_reader is None:
        _ocr_reader = easyocr.Reader(['en'], gpu=False, verbose=False)
    doc = fitz.open(file_path)
    full_text = ""
    # configurable OCR page limit: env OCR_MAX_PAGES (default 12), 0 = all pages
    try:
        env_max = int(os.getenv("OCR_MAX_PAGES", "12"))
    except:
        env_max = 12
    if env_max <= 0:
        max_pages = len(doc)
    else:
        max_pages = min(env_max, len(doc))
    print(f"OCR processing {max_pages}/{len(doc)} pages (OCR_MAX_PAGES={env_max})...")
    for i in range(max_pages):
        page = doc[i]
        pix = page.get_pixmap(dpi=100)
        img_bytes = pix.tobytes("png")
        nparr = np.frombuffer(img_bytes, np.uint8)
        img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
        result = _ocr_reader.readtext(img, detail=0, paragraph=True)
        full_text += "\n".join(result) + "\n"
    doc.close()
    return full_text.strip()

def _ocr_pdf_pages(file_path: str) -> list[dict]:
    """Per-page OCR, returns [{'page':1,'text':...}] using same OCR_MAX_PAGES limit."""
    global _ocr_reader
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return []
    try:
        import easyocr
        import numpy as np
        import cv2
    except ImportError:
        return []
    if _ocr_reader is None:
        _ocr_reader = easyocr.Reader(['en'], gpu=False, verbose=False)
    doc = fitz.open(file_path)
    try:
        env_max = int(os.getenv("OCR_MAX_PAGES", "12"))
    except:
        env_max = 12
    max_pages = len(doc) if env_max <= 0 else min(env_max, len(doc))
    pages = []
    for i in range(max_pages):
        page = doc[i]
        pix = page.get_pixmap(dpi=100)
        img_bytes = pix.tobytes("png")
        nparr = np.frombuffer(img_bytes, np.uint8)
        img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
        result = _ocr_reader.readtext(img, detail=0, paragraph=True)
        txt = _clean_text("\n".join(result).strip())
        if txt:
            pages.append({"page": i + 1, "text": txt})
    doc.close()
    return pages

def extract_text_from_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text += shape.text + "\n"
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
    return text.strip()

def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return _clean_text(f.read())

def _clean_text(text: str) -> str:
    # remove surrogates and invalid unicode (from OCR)
    return text.encode('utf-8', 'ignore').decode('utf-8', 'ignore')

def chunk_text(text: str, chunk_size: int = 350, overlap: int = 50) -> list[str]:
    text = _clean_text(text)
    # Sentence-aware chunking - avoids cutting in middle of concept
    # Split on sentence boundaries + bullet points
    sentences = re.split(r'(?<=[.!?])\s+|\n\s*•\s*|\n\s*-\s*', text)
    sentences = [re.sub(r"\s+", " ", s).strip() for s in sentences if s.strip()]
    if not sentences:
        words = text.split()
        sentences = [" ".join(words[i:i+50]) for i in range(0, len(words), 50)]

    chunks = []
    cur_words = []
    cur_len = 0
    for sent in sentences:
        w = sent.split()
        # if adding this sentence exceeds chunk_size, flush current chunk
        if cur_len + len(w) > chunk_size and cur_len > 100:
            chunk = " ".join(cur_words)
            if chunk:
                chunks.append(chunk)
            # overlap: keep last 'overlap' words
            if overlap > 0 and len(cur_words) > overlap:
                cur_words = cur_words[-overlap:] + w
                cur_len = len(cur_words)
            else:
                cur_words = w
                cur_len = len(w)
        else:
            cur_words.extend(w)
            cur_len += len(w)

    if cur_words:
        chunk = " ".join(cur_words)
        chunk = re.sub(r"\s+", " ", chunk).strip()
        if chunk:
            chunks.append(chunk)

    # Fallback: split any oversized chunks (PPTX bullets without periods often produce 1 huge chunk)
    final = []
    for ch in chunks:
        w = ch.split()
        if len(w) > int(chunk_size * 1.2):
            start = 0
            while start < len(w):
                end = start + chunk_size
                c = " ".join(w[start:end])
                if c:
                    final.append(c)
                if end >= len(w):
                    break
                start = end - overlap
        else:
            final.append(ch)
    # safety: if still single huge chunk (e.g. no sentence boundaries at all)
    if len(final) == 1 and len(final[0].split()) > chunk_size:
        w = final[0].split()
        final = []
        start = 0
        while start < len(w):
            end = start + chunk_size
            c = " ".join(w[start:end])
            if c:
                final.append(c)
            if end >= len(w):
                break
            start = end - overlap

    return final if final else chunks
